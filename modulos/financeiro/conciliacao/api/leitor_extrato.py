"""
Leitura de extrato bancário — OFX/OFC, CSV, XLSX.

Implementa references/conciliacao/formatos-extrato.md. A regra que governa
o arquivo inteiro é a da honestidade: arquivo que não pôde ser lido é dito
como tal, linha descartada é contada e informada, e nenhum lançamento é
deduzido de um total.

Ordem de preferência: OFX (tem FITID, que identifica reimportação com
certeza), depois CSV, depois XLSX. PDF não é processado aqui — PDF de
extrato costuma ser digitalizado e valor lido por OCR não é confiável.

Uso:
    resultado = ler_extrato(conteudo_bytes, "extrato.ofx")
    if resultado.precisa_mapeamento:
        ...  # perguntar ao usuário qual coluna é qual
    resultado.lancamentos  # list[Lancamento] pronto para o motor
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from .motor_conciliacao import Lancamento, normalizar_texto, parse_valor_br

# --------------------------------------------------------------------------
# Cabeçalhos conhecidos (references/conciliacao/formatos-extrato.md)
# --------------------------------------------------------------------------

CABECALHOS = {
    "data": ("data", "data lancamento", "data lancamento", "data movimento",
             "dt lanc", "data da movimentacao", "data mov"),
    "data_efetivacao": ("data efetivacao", "data efetiva", "data credito",
                        "data compensacao", "data liquidacao"),
    "descricao": ("historico", "descricao", "lancamento", "memo",
                  "historico complementar", "detalhamento"),
    "valor": ("valor", "valor r", "montante", "vlr", "valor lancamento"),
    "tipo": ("tipo", "d c", "dc", "natureza", "debito credito",
             "tipo lancamento"),
    "documento": ("documento", "doc", "n documento", "numero documento",
                  "cpf cnpj", "cpf cnpj contraparte", "num doc"),
    "entrada": ("entrada", "credito", "creditos", "valor credito"),
    "saida": ("saida", "debito", "debitos", "valor debito"),
    "contraparte": ("contraparte", "favorecido", "beneficiario", "pagador",
                    "nome contraparte", "cliente fornecedor"),
    "saldo": ("saldo", "saldo apos", "saldo atual", "saldo do dia"),
}

# Linhas que não são lançamento e aparecem em quase todo extrato.
PALAVRAS_NAO_LANCAMENTO = (
    "saldo anterior", "saldo final", "saldo do dia", "saldo em conta",
    "total", "totalizador", "extrato de", "periodo", "agencia conta",
    "s a l d o",
)

FORMATOS_DATA = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%d.%m.%Y",
                 "%Y/%m/%d", "%Y%m%d")


# --------------------------------------------------------------------------
# Resultado
# --------------------------------------------------------------------------

@dataclass
class ResultadoLeitura:
    lancamentos: list[Lancamento] = field(default_factory=list)
    formato: str = ""
    encoding: str = ""
    ignoradas: int = 0
    motivos_ignoradas: dict[str, int] = field(default_factory=dict)
    avisos: list[str] = field(default_factory=list)
    erro: str = ""
    precisa_mapeamento: bool = False
    colunas: list[str] = field(default_factory=list)
    amostra: list[list[str]] = field(default_factory=list)
    conta_mascarada: str = ""

    @property
    def ok(self) -> bool:
        return not self.erro and not self.precisa_mapeamento

    def resumo(self) -> str:
        """Frase verificável: quantos lançamentos e quantas linhas descartadas."""
        if self.erro:
            return self.erro
        if self.precisa_mapeamento:
            return ("Não reconheci o cabeçalho deste arquivo. "
                    "Informe qual coluna é a data, o valor e o histórico.")
        partes = [f"Li {len(self.lancamentos)} lançamento(s)"]
        if self.ignoradas:
            detalhe = ", ".join(f"{qtd} {motivo}"
                                for motivo, qtd in sorted(self.motivos_ignoradas.items()))
            partes.append(f"ignorei {self.ignoradas} linha(s) ({detalhe})")
        return " e ".join(partes) + "."

    def _ignorar(self, motivo: str) -> None:
        self.ignoradas += 1
        self.motivos_ignoradas[motivo] = self.motivos_ignoradas.get(motivo, 0) + 1


# --------------------------------------------------------------------------
# Entrada principal
# --------------------------------------------------------------------------

def ler_extrato(conteudo: bytes, nome_arquivo: str = "",
                mapeamento: dict[str, str] | None = None) -> ResultadoLeitura:
    """Lê o arquivo e devolve lançamentos + contagem do que foi descartado."""
    nome = (nome_arquivo or "").lower()

    if not conteudo:
        return ResultadoLeitura(erro="Arquivo vazio — não há o que processar.")

    if len(conteudo) > 10 * 1024 * 1024:
        aviso = ("Arquivo acima de 10 MB. A leitura pode demorar; "
                 "se possível, exporte o extrato em períodos menores.")
    else:
        aviso = ""

    if nome.endswith(".pdf") or conteudo[:5] == b"%PDF-":
        return ResultadoLeitura(
            formato="pdf",
            erro=("Não processo extrato em PDF: quando o arquivo é digitalizado, "
                  "o valor lido por OCR não é confiável. Exporte o extrato em "
                  "OFX ou CSV pelo internet banking e eu processo."),
        )

    if nome.endswith((".xlsx", ".xlsm")) or conteudo[:2] == b"PK":
        resultado = _ler_xlsx(conteudo, mapeamento)
    elif nome.endswith((".ofx", ".ofc")) or _parece_ofx(conteudo):
        resultado = _ler_ofx(conteudo)
    else:
        resultado = _ler_csv(conteudo, mapeamento)

    if aviso:
        resultado.avisos.insert(0, aviso)
    if resultado.ok and not resultado.lancamentos:
        resultado.avisos.append(
            "Nenhum lançamento foi encontrado no arquivo — confira se o "
            "período exportado tem movimentação."
        )
    return resultado


# --------------------------------------------------------------------------
# OFX
# --------------------------------------------------------------------------

def _parece_ofx(conteudo: bytes) -> bool:
    inicio = conteudo[:2048].upper()
    return b"OFXHEADER" in inicio or b"<OFX>" in inicio or b"<STMTTRN>" in inicio


def _ler_ofx(conteudo: bytes) -> ResultadoLeitura:
    texto, encoding = _decodificar(conteudo)
    resultado = ResultadoLeitura(formato="ofx", encoding=encoding)

    conta = re.search(r"<ACCTID>([^<\r\n]+)", texto, re.IGNORECASE)
    if conta:
        resultado.conta_mascarada = mascarar_conta(conta.group(1))

    blocos = re.findall(r"<STMTTRN>(.*?)</STMTTRN>", texto,
                        re.IGNORECASE | re.DOTALL)
    if not blocos:
        resultado.erro = ("Arquivo OFX sem transações (<STMTTRN>). "
                          "Confira se a exportação cobriu o período desejado.")
        return resultado

    for i, bloco in enumerate(blocos, start=1):
        def campo(tag: str) -> str:
            m = re.search(rf"<{tag}>([^<\r\n]*)", bloco, re.IGNORECASE)
            return (m.group(1) or "").strip() if m else ""

        bruto_data = campo("DTPOSTED")[:8]
        try:
            data_lanc = datetime.strptime(bruto_data, "%Y%m%d").date()
        except ValueError:
            resultado._ignorar("sem data válida")
            continue

        try:
            centavos = parse_valor_br(campo("TRNAMT"))
        except ValueError:
            resultado._ignorar("valor não numérico")
            continue

        # TRNTYPE DEBIT com valor positivo acontece em alguns bancos;
        # o sinal do tipo manda, porque somar crédito com débito produz
        # um painel plausível e errado.
        tipo = campo("TRNTYPE").upper()
        if tipo in ("DEBIT", "PAYMENT", "FEE", "SRVCHG") and centavos > 0:
            centavos = -centavos
        elif tipo in ("CREDIT", "DEP", "DIRECTDEP") and centavos < 0:
            centavos = abs(centavos)

        fitid = campo("FITID")
        descricao = " ".join(p for p in (campo("MEMO"), campo("NAME")) if p).strip()

        resultado.lancamentos.append(Lancamento(
            id=fitid or f"OFX{i:05d}",
            data=data_lanc,
            descricao=descricao,
            valor_centavos=centavos,
            documento=campo("CHECKNUM") or campo("REFNUM"),
            contraparte=campo("NAME"),
            fitid=fitid,
        ))

    if not any(l.fitid for l in resultado.lancamentos):
        resultado.avisos.append(
            "O OFX veio sem FITID. A detecção de reimportação cai para a "
            "chave (data, valor, histórico), que é heurística."
        )
    return resultado


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------

def _ler_csv(conteudo: bytes, mapeamento: dict[str, str] | None) -> ResultadoLeitura:
    texto, encoding = _decodificar(conteudo)
    resultado = ResultadoLeitura(formato="csv", encoding=encoding)

    linhas = [l for l in texto.splitlines() if l.strip()]
    if not linhas:
        resultado.erro = "Arquivo sem conteúdo legível."
        return resultado

    separador = _detectar_separador(linhas)
    tabela = list(csv.reader(io.StringIO("\n".join(linhas)), delimiter=separador))
    return _montar_de_tabela(tabela, resultado, mapeamento)


def _detectar_separador(linhas: list[str]) -> str:
    """`;` é o mais comum no Brasil, mas quem decide é a contagem."""
    amostra = "\n".join(linhas[:20])
    contagens = {sep: amostra.count(sep) for sep in (";", ",", "\t", "|")}
    melhor = max(contagens, key=lambda s: contagens[s])
    return melhor if contagens[melhor] else ";"


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------

def _ler_xlsx(conteudo: bytes, mapeamento: dict[str, str] | None) -> ResultadoLeitura:
    resultado = ResultadoLeitura(formato="xlsx", encoding="—")
    try:
        from openpyxl import load_workbook
    except ImportError:
        resultado.erro = (
            "Não consigo ler XLSX neste ambiente: a biblioteca openpyxl não "
            "está instalada. Instale-a (pip install openpyxl) ou exporte o "
            "extrato em CSV ou OFX."
        )
        return resultado

    try:
        planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as erro:            # arquivo corrompido, protegido por senha…
        resultado.erro = f"Não consegui abrir a planilha: {erro}"
        return resultado

    aba = planilha[planilha.sheetnames[0]]
    if len(planilha.sheetnames) > 1:
        resultado.avisos.append(
            f"A planilha tem {len(planilha.sheetnames)} abas; li a primeira "
            f"('{planilha.sheetnames[0]}')."
        )

    tabela: list[list[str]] = []
    for linha in aba.iter_rows(values_only=True):
        tabela.append(["" if c is None else str(c).strip() for c in linha])
    planilha.close()

    return _montar_de_tabela(tabela, resultado, mapeamento)


# --------------------------------------------------------------------------
# Tabela → lançamentos (comum a CSV e XLSX)
# --------------------------------------------------------------------------

def _montar_de_tabela(tabela: list[list[str]], resultado: ResultadoLeitura,
                      mapeamento: dict[str, str] | None) -> ResultadoLeitura:
    linha_cabecalho, indices = _localizar_cabecalho(tabela)

    if mapeamento:
        indices = {campo: int(pos) for campo, pos in mapeamento.items()
                   if str(pos).strip() != ""}
        linha_cabecalho = linha_cabecalho if linha_cabecalho >= 0 else 0

    tem_valor = "valor" in indices or ("entrada" in indices or "saida" in indices)
    if not ("data" in indices or "data_efetivacao" in indices) or not tem_valor:
        # Perguntar custa uma mensagem; adivinhar errado custa a conciliação.
        resultado.precisa_mapeamento = True
        base = max(linha_cabecalho, 0)
        resultado.colunas = [c for c in (tabela[base] if base < len(tabela) else [])]
        resultado.amostra = tabela[base:base + 6]
        return resultado

    for bruta in tabela[linha_cabecalho + 1:]:
        if not any(str(c).strip() for c in bruta):
            continue

        texto_linha = normalizar_texto(" ".join(str(c) for c in bruta))
        if any(p in texto_linha for p in PALAVRAS_NAO_LANCAMENTO) and \
                not _tem_data(bruta, indices):
            resultado._ignorar("de cabeçalho/saldo")
            continue

        data_lanc = _extrair_data(bruta, indices)
        if data_lanc is None:
            resultado._ignorar("sem data válida")
            continue

        try:
            centavos = _extrair_valor(bruta, indices)
        except ValueError:
            resultado._ignorar("valor não numérico")
            continue

        descricao = _celula(bruta, indices.get("descricao"))
        contraparte = _celula(bruta, indices.get("contraparte"))
        documento = _celula(bruta, indices.get("documento"))

        lanc = Lancamento(
            id=f"L{len(resultado.lancamentos) + 1:05d}",
            data=data_lanc,
            descricao=descricao,
            valor_centavos=centavos,
            documento=documento,
            contraparte=contraparte,
        )
        resultado.lancamentos.append(lanc)

        if data_lanc > date.today():
            resultado.avisos.append(
                f"Lançamento {lanc.id} tem data futura ({data_lanc:%d/%m/%Y})."
            )
        if centavos == 0:
            resultado.avisos.append(f"Lançamento {lanc.id} tem valor zero.")

    if "descricao" not in indices:
        resultado.avisos.append(
            "O arquivo não traz histórico/descrição. Sem ele, só as regras de "
            "documento e de valor exato conseguem casar lançamentos."
        )
    if "data_efetivacao" in indices and "data" in indices:
        resultado.avisos.append(
            "O arquivo traz data de lançamento e de efetivação; usei a de "
            "efetivação para o cruzamento."
        )
    return resultado


def _localizar_cabecalho(tabela: list[list[str]]) -> tuple[int, dict[str, int]]:
    """Procura nas primeiras linhas a que parece cabeçalho.

    Extratos trazem linhas institucionais antes do cabeçalho real, então
    não dá para assumir que ele é a primeira linha.
    """
    melhor_linha, melhores_indices = -1, {}
    for i, linha in enumerate(tabela[:15]):
        indices: dict[str, int] = {}
        for j, celula in enumerate(linha):
            campo = _campo_do_cabecalho(str(celula))
            if campo and campo not in indices:
                indices[campo] = j
        if len(indices) > len(melhores_indices):
            melhor_linha, melhores_indices = i, indices
    return melhor_linha, melhores_indices


def _campo_do_cabecalho(texto: str) -> str:
    alvo = normalizar_texto(texto)
    if not alvo:
        return ""
    for campo, variacoes in CABECALHOS.items():
        if alvo in variacoes:
            return campo
    for campo, variacoes in CABECALHOS.items():
        if any(alvo.startswith(v) or v.startswith(alvo) for v in variacoes if v):
            return campo
    return ""


def _celula(linha: list[str], indice: int | None) -> str:
    if indice is None or indice >= len(linha):
        return ""
    return str(linha[indice]).strip()


def _tem_data(linha: list[str], indices: dict[str, int]) -> bool:
    return _extrair_data(linha, indices) is not None


def _extrair_data(linha: list[str], indices: dict[str, int]) -> date | None:
    """Prefere a data de efetivação: é ela que casa com o extrato do banco."""
    for campo in ("data_efetivacao", "data"):
        bruto = _celula(linha, indices.get(campo))
        if not bruto:
            continue
        convertida = _converter_data(bruto)
        if convertida:
            return convertida
    return None


def _converter_data(bruto: str) -> date | None:
    bruto = bruto.strip()
    if not bruto:
        return None
    # Célula de planilha já convertida para datetime pelo openpyxl.
    if re.match(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$", bruto):
        bruto = bruto[:10]
    for formato in FORMATOS_DATA:
        try:
            return datetime.strptime(bruto, formato).date()
        except ValueError:
            continue
    return None


def _extrair_valor(linha: list[str], indices: dict[str, int]) -> int:
    """Resolve as três convenções de sinal antes de somar qualquer coisa."""
    # Convenção 3 — colunas separadas de entrada e saída.
    if "entrada" in indices or "saida" in indices:
        entrada = _celula(linha, indices.get("entrada"))
        saida = _celula(linha, indices.get("saida"))
        if entrada and _numerico(entrada):
            return abs(parse_valor_br(entrada))
        if saida and _numerico(saida):
            return -abs(parse_valor_br(saida))
        raise ValueError("sem valor de entrada ou saída")

    bruto = _celula(linha, indices.get("valor"))
    if not _numerico(bruto):
        raise ValueError("valor não numérico")
    centavos = parse_valor_br(bruto)

    # Convenção 1 — coluna de tipo C/D manda no sinal.
    tipo = normalizar_texto(_celula(linha, indices.get("tipo")))
    if tipo:
        if tipo.startswith("d") or "debito" in tipo or "saida" in tipo:
            return -abs(centavos)
        if tipo.startswith("c") or "credito" in tipo or "entrada" in tipo:
            return abs(centavos)

    # Convenção 2 — o próprio valor já vem com sinal.
    return centavos


def _numerico(texto: str) -> bool:
    return bool(re.search(r"\d", texto or ""))


# --------------------------------------------------------------------------
# Encoding, máscara e reimportação
# --------------------------------------------------------------------------

def _decodificar(conteudo: bytes) -> tuple[str, str]:
    """UTF-8 primeiro; bancos brasileiros exportam muito em latin-1."""
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return conteudo.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return conteudo.decode("latin-1", errors="replace"), "latin-1 (com perdas)"


def mascarar_conta(conta: str) -> str:
    """`****1234` — conta e agência não aparecem inteiras na tela nem no log."""
    digitos = re.sub(r"\D", "", conta or "")
    return f"****{digitos[-4:]}" if len(digitos) >= 4 else "****"


def chave_reimportacao(lanc: Lancamento) -> str:
    """FITID quando houver; senão (data, valor, histórico normalizado)."""
    if lanc.fitid:
        return f"fitid:{lanc.fitid}"
    return f"dvh:{lanc.data.isoformat()}|{lanc.valor_centavos}|{normalizar_texto(lanc.descricao)}"


def detectar_reimportacao(novos: list[Lancamento],
                          ja_importados: list[Lancamento]) -> list[Lancamento]:
    """Lançamentos que já entraram antes. Reimportar sem perguntar concilia
    duas vezes o mesmo título."""
    conhecidos = {chave_reimportacao(l) for l in ja_importados}
    return [l for l in novos if chave_reimportacao(l) in conhecidos]
